"""Gera os CSVs de carga a partir da planilha da curadoria V2 (HU13 e HU15).

Uso (na raiz do repositório):
    python3 docs/curadoria/carga/gerar_carga.py

Ferramenta de documentação, fora do produto: só lê a planilha assinada e escreve os CSVs que a
curadoria importa pelas telas (prévia, rascunho e aprovação por outro curador). Nenhum dado entra
em vigor por este script. Requer openpyxl.
"""

import csv
import datetime
import re
import unicodedata
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[3]
PLANILHA = RAIZ / 'docs/curadoria/hu14-pendencias-do-curador-v2-revisado.xlsx'
SAIDA = RAIZ / 'docs/curadoria/carga'
LINHAS_POR_ARQUIVO = 1000

COLUNAS_TABELA = [
    'tipoAto', 'nomeDoAto', 'ramo', 'dias', 'unidade', 'contagem', 'politicaDobro', 'calculo',
    'marcoInicial', 'localDoAto', 'alerta', 'fundamento', 'fonteUrl', 'vigenciaInicio',
    'vigenciaFim',
]
COLUNAS_CALENDARIO = [
    'abrangencia', 'uf', 'municipioIbge', 'tribunal', 'comarca', 'tipo', 'inicio', 'fim',
    'descricao', 'atoNormativo', 'urlAto',
]

RAMOS = {
    'cível': 'civel', 'trabalhista': 'trabalhista', 'penal': 'penal',
    'juizado cível': 'juizados', 'juizado criminal': 'juizado-criminal',
}
CONTAGENS = {'úteis': 'uteis', 'corridos': 'corridos', 'horas corridas': 'corridos'}
DOBRO = [
    ('Em dobro para Fazenda, MP e Defensoria', 'entes-cpc'),
    ('Em dobro só para a Defensoria', 'so-defensoria'),
    ('Fazenda em dobro para recorrer', 'fazenda-para-recorrer'),
    ('Não dobra', 'nao-dobra'),
    ('Sem dobro automático', 'verificar'),
]
LOCAIS = {
    'Órgão onde o processo tramita': 'orgao-do-processo',
    '1º grau (juízo de origem)': 'juizo-de-origem',
    'Tribunal de 2º grau (tribunal recorrido)': 'tribunal-de-segundo-grau',
    'Órgão que proferiu a decisão': 'orgao-da-decisao',
    'Tribunal superior': 'tribunal-superior',
    'Não se aplica': 'nao-se-aplica',
}
# Tipos do calendário (HU13). "Expediente reduzido" tem o efeito da indisponibilidade no motor: só
# prorroga início e vencimento (CPC, art. 224, §1º); a descrição mantém o texto da curadora.
TIPOS_DE_EVENTO = {
    'feriado': 'feriado', 'recesso': 'recesso', 'suspensão de prazos': 'portaria',
    'indisponibilidade': 'indisponibilidade', 'expediente reduzido': 'indisponibilidade',
}


def codigo(nome: str) -> str:
    sem_acento = unicodedata.normalize('NFKD', nome).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', sem_acento.lower()).strip('-')[:80]


def calculo(texto: str) -> str:
    if texto.startswith('Não calcula'):
        return 'nao-calcula'
    if texto.startswith('Com data informada') or 'encadeado' in texto:
        return 'data-informada'
    return 'automatico'


def data_iso(valor) -> str:
    if isinstance(valor, datetime.datetime):
        return valor.date().isoformat()
    d, m, a = str(valor).strip().split('/')
    return f'{a}-{m}-{d}'


def escrever(nome: str, colunas: list[str], linhas: list[dict]) -> None:
    with open(SAIDA / nome, 'w', newline='', encoding='utf-8') as arquivo:
        saida = csv.DictWriter(arquivo, fieldnames=colunas, delimiter=';', lineterminator='\n')
        saida.writeheader()
        saida.writerows(linhas)


def tabela(livro) -> tuple[list[dict], list[str]]:
    aba4 = {str(r[0]): r for r in livro['4. Tabela de prazos'].iter_rows(min_row=2, values_only=True) if r[1]}
    linhas, fora = [], []
    for r in livro['9. Tabela do motor'].iter_rows(min_row=2, values_only=True):
        if not r[2] or r[1] != 'Ativo':
            continue
        numero, _, ato, ramo, prazo, unidade, contagem, marco, calc, _, dobro, local, fund, alerta = r[:14]
        if not isinstance(prazo, (int, float)):
            fora.append(f'linha {numero} ({ato}, {ramo}): sem prazo numérico ("{calc}")')
            continue
        origem = aba4[str(numero)]
        url = re.findall(r'https://\S+', str(origem[8]))[0]
        politica = next(valor for prefixo, valor in DOBRO if str(dobro).startswith(prefixo))
        linhas.append({
            'tipoAto': codigo(ato), 'nomeDoAto': ato, 'ramo': RAMOS[ramo], 'dias': int(prazo),
            'unidade': unidade, 'contagem': CONTAGENS[contagem], 'politicaDobro': politica,
            'calculo': calculo(str(calc)), 'marcoInicial': (marco or '').strip(),
            'localDoAto': LOCAIS.get(local, ''), 'alerta': (alerta or '').strip(),
            'fundamento': str(fund).strip(), 'fonteUrl': url,
            'vigenciaInicio': data_iso(origem[9]), 'vigenciaFim': '',
        })
    return linhas, fora


def calendario(livro) -> tuple[list[dict], dict[str, int]]:
    linhas, fora = [], {}
    for r in livro['5. Calendário'].iter_rows(min_row=11, values_only=True):
        abrangencia, nome, _, tipo, periodo, descricao, ato, url = r[:8]
        if not tipo or tipo in ('Tipo', 'cadastro territorial'):
            continue
        nome = str(nome)
        if abrangencia == 'nacional' and tipo == 'suspensão de prazos':
            # RN02 (20/12 a 20/01) é regra do motor, por ramo: não entra no calendário.
            fora['suspensão nacional de 20/12 a 20/01 (regra do motor, RN02)'] = fora.get('suspensão nacional de 20/12 a 20/01 (regra do motor, RN02)', 0) + 1
            continue
        if abrangencia not in ('nacional', 'tribunal') or '·' in nome:
            chave = {
                'UF': 'feriado estadual publicado por um tribunal (modelo sem "tribunal + UF")',
                'município': 'feriado municipal publicado por um tribunal (modelo sem "tribunal + município")',
                'tribunal': 'TRF2 2º grau separado das seções (modelo sem calendário por grau)',
                'comarca': 'linha de exemplo da planilha (fictícia)',
            }[abrangencia]
            fora[chave] = fora.get(chave, 0) + 1
            continue
        inicio, _, fim = str(periodo).partition(' a ')
        sigla = '' if abrangencia == 'nacional' else re.match(r'[A-Z]+-?\d*', nome).group().replace('-', '')
        linhas.append({
            'abrangencia': abrangencia, 'uf': '', 'municipioIbge': '', 'tribunal': sigla,
            'comarca': '', 'tipo': TIPOS_DE_EVENTO[tipo], 'inicio': data_iso(inicio),
            'fim': data_iso(fim or inicio), 'descricao': str(descricao).strip()[:500],
            'atoNormativo': str(ato).strip()[:500], 'urlAto': str(url).strip().split()[0],
        })
    return linhas, fora


def main() -> None:
    livro = openpyxl.load_workbook(PLANILHA, data_only=True)
    linhas_tabela, fora_tabela = tabela(livro)
    escrever('tabela-de-prazos.csv', COLUNAS_TABELA, linhas_tabela)
    linhas_cal, fora_cal = calendario(livro)
    for i in range(0, len(linhas_cal), LINHAS_POR_ARQUIVO):
        escrever(f'calendario-{i // LINHAS_POR_ARQUIVO + 1:02d}.csv', COLUNAS_CALENDARIO, linhas_cal[i:i + LINHAS_POR_ARQUIVO])
    print(f'tabela: {len(linhas_tabela)} linhas; fora: {len(fora_tabela)}')
    for item in fora_tabela:
        print(f'  - {item}')
    print(f'calendário: {len(linhas_cal)} linhas')
    for motivo, total in fora_cal.items():
        print(f'  - fora: {total} · {motivo}')


if __name__ == '__main__':
    main()
